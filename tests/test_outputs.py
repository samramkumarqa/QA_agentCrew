import io
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import SAMPLE
from fakes import ScriptedLLM
from qacrew.crew import run_crew
from qacrew.documents import DocumentError, read_document
from qacrew.excel import build_workbook

PROJECT2 = Path(__file__).resolve().parents[2] / "QA_generateAutomationcodes"


@pytest.fixture(scope="module")
def result():
    return run_crew(SAMPLE.read_text(), ScriptedLLM(), project="Swag Labs")


def test_workbook_sheets_and_approval(result):
    data = build_workbook(result, "Swag Labs", "test-model", approved={"TC-001", "TC-004"})
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Test Cases", "Requirements", "Traceability", "Review Findings",
                             "BA Questions", "Automation Plan"]
    rows = list(wb["Test Cases"].iter_rows(values_only=True))
    assert rows[0][:9] == ("Test Case ID", "Requirement ID", "Title", "Type", "Priority", "Preconditions", "Steps",
                           "Test Data", "Expected Result")
    assert [r[0] for r in rows[1:]] == ["TC-001", "TC-004"]
    assert rows[2][6].startswith("1. Add a product\n2. Open the cart") and rows[2][9] == "API · High"
    summary = dict(wb["Summary"].iter_rows(values_only=True))
    assert summary["Requirement coverage (draft → after review)"] == "67% → 100%"
    trace = list(wb["Traceability"].iter_rows(values_only=True))
    assert trace[3][4] == "Covered" and trace[3][2] == "TC-004"


@pytest.mark.skipif(not (PROJECT2 / "pwgen").exists(), reason="AI Automation Code Generator not checked out next to this repo")
def test_project2_reads_the_test_pack(result):
    sys.path.insert(0, str(PROJECT2))
    from pwgen.testcases import read_test_cases
    cases = read_test_cases(build_workbook(result, "Swag Labs", "m"))
    assert [c.id for c in cases] == ["TC-001", "TC-002", "TC-003", "TC-004", "TC-005"]
    assert cases[3].requirement_id == "REQ-03" and cases[3].steps[-1] == "Click Continue"


def test_read_documents(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Users must log in with a username and a password before shopping.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Tax", "8%"
    buf = io.BytesIO()
    d.save(buf)
    text = read_document("brd.docx", buf.getvalue())
    assert "log in" in text and "Tax | 8%" in text
    assert "Tax is 8%" in read_document("brd.md", SAMPLE.read_bytes())
    with pytest.raises(DocumentError, match="PDF, Word"):
        read_document("brd.xlsx", b"x" * 100)
    with pytest.raises(DocumentError, match="no readable text"):
        read_document("empty.txt", b"  ")
