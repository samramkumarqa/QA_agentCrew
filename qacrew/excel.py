"""The Excel test pack. The "Test Cases" sheet uses the same columns as the AI Test Case Generator,
so the AI Automation Code Generator can read it directly."""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .checks import traceability
from .crew import CrewResult

HEAD = PatternFill("solid", fgColor="14213D")
GOOD = PatternFill("solid", fgColor="DFF0E3")
BAD = PatternFill("solid", fgColor="FCE3CF")


def _sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    return ws


def build_workbook(result: CrewResult, project: str, model: str, approved: set[str] | None = None) -> bytes:
    final = [c for c in result.final.test_cases if approved is None or c.id in approved]
    plan = {i.test_case_id: i for i in result.automation.items}
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    rows = [
        ("Project", project), ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")), ("Model", model),
        ("Requirements", len(result.requirements.requirements)),
        ("Questions for the BA", len(result.requirements.questions)),
        ("Test cases (draft → after review)", f"{len(result.draft.test_cases)} → {len(result.final.test_cases)}"),
        ("Test cases in this file (approved)", len(final)),
        ("Requirement coverage (draft → after review)",
         f"{result.coverage_before.percent}% → {result.coverage_after.percent}%"),
        ("Review findings", len(result.review.findings)),
        ("Revision", f"added {len(result.changes['added'])}, updated {len(result.changes['updated'])}, "
                     f"removed {len(result.changes['removed'])}"),
        ("Review summary", result.review.summary),
        ("Automation summary", result.automation.summary),
    ]
    if result.truncated:
        rows.append(("Note", "The document was long; only its first part was analysed."))
    for k, v in rows:
        ws.append([k, v])
    for c in ws["A"]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 90
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    _sheet(wb, "Test Cases",
           ["Test Case ID", "Requirement ID", "Title", "Type", "Priority", "Preconditions", "Steps", "Test Data",
            "Expected Result", "Automation"],
           [[c.id, ", ".join(c.requirement_ids), c.title, c.type, c.priority, c.preconditions,
             "\n".join(f"{i}. {s}" for i, s in enumerate(c.steps, 1)), c.test_data, c.expected_result,
             (f"{plan[c.id].layer} · {plan[c.id].priority}" if c.id in plan else "")] for c in final],
           [12, 14, 36, 11, 10, 24, 50, 24, 40, 16])

    _sheet(wb, "Requirements", ["ID", "Title", "Description", "Acceptance Criteria", "Priority"],
           [[r.id, r.title, r.description, "\n".join(f"- {a}" for a in r.acceptance_criteria), r.priority]
            for r in result.requirements.requirements], [10, 28, 50, 60, 10])

    tw = _sheet(wb, "Traceability", ["Requirement", "Title", "Test Cases", "Types", "Status"],
                [[t["requirement"], t["title"], ", ".join(t["test_cases"]), ", ".join(t["types"]), t["status"]]
                 for t in traceability(result.requirements, result.final)], [12, 32, 40, 28, 14])
    for row in tw.iter_rows(min_row=2):
        row[4].fill = GOOD if row[4].value == "Covered" else BAD

    _sheet(wb, "Review Findings", ["Severity", "Kind", "Target", "Finding", "Suggestion"],
           [[f.severity, f.kind.replace("_", " "), f.target, f.detail, f.suggestion] for f in result.review.findings],
           [10, 18, 12, 60, 60])
    ws = wb["Review Findings"]
    ws.append([])
    ws.append(["Revision", "", "", result.revision.notes, ""])

    _sheet(wb, "BA Questions", ["Requirement", "Question", "Why it matters"],
           [[q.requirement_id, q.question, q.why_it_matters] for q in result.requirements.questions], [12, 70, 60])

    titles = {c.id: c.title for c in result.final.test_cases}
    _sheet(wb, "Automation Plan", ["Test Case ID", "Title", "Priority", "Layer", "Reason"],
           [[i.test_case_id, titles.get(i.test_case_id, ""), i.priority, i.layer, i.reason]
            for i in sorted(result.automation.items, key=lambda i: ("High", "Medium", "Low").index(i.priority))],
           [12, 36, 10, 14, 70])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
