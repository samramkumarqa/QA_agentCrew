"""Read the BRD from PDF, Word, Markdown or text."""

from __future__ import annotations

import io


class DocumentError(ValueError):
    pass


def read_document(name: str, data: bytes) -> str:
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    try:
        if ext == "pdf":
            import pdfplumber

            with pdfplumber.open(io.BytesIO(data)) as pdf:
                text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        elif ext == "docx":
            import docx

            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        elif ext in {"txt", "md", "markdown"}:
            text = data.decode("utf-8", errors="replace")
        else:
            raise DocumentError("Upload a PDF, Word (.docx), Markdown or text file.")
    except DocumentError:
        raise
    except Exception as e:
        raise DocumentError(f"Couldn't read {name}: {e}") from e
    text = text.strip()
    if len(text) < 40:
        raise DocumentError(f"{name} has no readable text (a scanned PDF needs OCR first).")
    return text
